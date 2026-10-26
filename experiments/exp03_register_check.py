"""Exp 03: human check of the blind clinical register labels (D2) on the pilot.

40 pilot items drawn with a fixed seed, stratified by the labeler's class (14 technical,
13 lay, 13 both), go to a spreadsheet that shows only the vignette, the generator's phrase
and the HPO concept, never the labeler's answer. The author fills two yes/no columns
(technical term present? lay wording present?). Scoring gives Cohen's kappa between author
and labeler on the 3 classes; below KAPPA_MIN the D2 strata are reported as exploratory only
(PREREG, Desvios).

    uv run python experiments/exp03_register_check.py --make
    uv run python experiments/exp03_register_check.py --score reports/review/exp03_register_check_40.xlsx
"""
# ruff: noqa: E501  (PT-BR instruction text below is prose)

from __future__ import annotations

import argparse
import json

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from hpolab.casemetrics import cohen_kappa
from hpolab.data import RAW_DIR, ROOT
from hpolab.hpo import load_babelon, load_ontology, pt_labels
from hpolab.llm import Prompt
from hpolab.vignettes import GENERATOR_MODEL, GENERATOR_PROMPT, final_vignettes

SEED = 42
PER_CLASS = {"technical": 14, "lay": 13, "both": 13}
KAPPA_MIN = 0.6
CLASSES = ("technical", "lay", "both")
PILOT = ROOT / "results" / "exp03" / "pilot"
LABELS = PILOT / "register_check.csv"
KEY = PILOT / "register_human_key.csv"
RESULT = PILOT / "register_kappa.json"
SHEET = ROOT / "reports" / "review" / "exp03_register_check_40.xlsx"
YES = {"sim", "s", "yes", "y"}
NO = {"não", "nao", "n", "no"}

INSTRUCTIONS = [
    "Conferência às cegas do registro dos itens (Exp 03, piloto)",
    "",
    "Cada linha é um achado do gabarito de uma vinheta sintética. A coluna 'trecho' é a frase que o gerador indicou para o achado; 'conceito HPO' diz qual achado procurar. A classificação do outro rotulador não aparece aqui.",
    "",
    "Leia a vinheta INTEIRA: o mesmo achado pode aparecer duas vezes (por exemplo, em palavras da família na história e com o termo médico no exame).",
    "",
    "Coluna 'termo técnico?': sim se o achado aparece em algum lugar da vinheta com termo médico, isto é, o nome que um médico usaria no prontuário ou num artigo, inclusive siglas (CIV, TDAH), tenha ou não tradução oficial no HPO. Exemplos: 'atraso global do desenvolvimento', 'plaquetopenia', 'derrame pleural', 'hipermobilidade articular', 'perímetro cefálico abaixo de -2 DP'.",
    "",
    "Coluna 'descrição leiga?': sim se o achado aparece (também ou só) com palavras do dia a dia ou uma descrição simples, sem o termo médico. Exemplos: 'demorou para sentar, andar e falar', 'fica com manchas roxas fácil', 'a pele parece sobrando', 'olhos bem afastados'.",
    "",
    "Preencha as duas colunas com sim ou não. Se não achar o achado na vinheta, marque não/não e explique em 'observação'.",
    "",
    "Quando terminar, salve o arquivo com o mesmo nome e avise. O script calcula o kappa de Cohen entre você e o rotulador.",
]


def pilot_texts() -> dict[str, str]:
    sha = Prompt.load(ROOT / "prompts" / GENERATOR_PROMPT).sha256
    with (ROOT / "results" / "exp03" / "vignettes.jsonl").open(encoding="utf-8") as fh:
        recs = [json.loads(line) for line in fh if line.strip()]
    return {
        cid: r["vignette"]
        for cid, r in final_vignettes(recs, "pilot", sha, GENERATOR_MODEL).items()
    }


def draw_items(labels: pd.DataFrame) -> pd.DataFrame:
    rng = np.random.default_rng(SEED)
    parts = []
    for cls in CLASSES:
        pool = labels[labels["man_d2"] == cls].sort_values(["case_id", "hpo_id"])
        pick = rng.choice(len(pool), PER_CLASS[cls], replace=False)
        parts.append(pool.iloc[np.sort(pick)])
    sample = pd.concat(parts)
    sample = sample.iloc[rng.permutation(len(sample))].reset_index(drop=True)
    sample.insert(0, "id", range(1, len(sample) + 1))
    return sample


def make() -> None:
    labels = pd.read_csv(LABELS, dtype=str, keep_default_na=False)
    texts = pilot_texts()
    ont = load_ontology(RAW_DIR / "hp.json")
    pt = pt_labels(load_babelon(RAW_DIR / "hp-pt.babelon.tsv"), ont).set_index("hpo_id")["pt_label"]
    sample = draw_items(labels)
    sample[["id", "case_id", "hpo_id", "man_d2"]].rename(columns={"man_d2": "labeler_d2"}).to_csv(
        KEY, index=False
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "itens"
    head = [
        "id",
        "vinheta",
        "trecho",
        "conceito HPO",
        "termo técnico?",
        "descrição leiga?",
        "observação",
    ]
    ws.append(head)
    for r in sample.itertuples():
        concept = r.label_en + (f" / {pt[r.hpo_id]}" if r.hpo_id in pt.index else "")
        ws.append([r.id, texts[r.case_id], r.phrase, concept, "", "", ""])
    widths = {"A": 5, "B": 90, "C": 35, "D": 30, "E": 14, "F": 14, "G": 30}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        for cell in row[4:7]:
            cell.fill = fill
    dv = DataValidation(type="list", formula1='"sim,não"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:F{len(sample) + 1}")
    ws.freeze_panes = "B2"
    info = wb.create_sheet("instruções", 0)
    for line in INSTRUCTIONS:
        info.append([line])
    info.column_dimensions["A"].width = 120
    for row in info.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    info["A1"].font = Font(bold=True, size=13)
    SHEET.parent.mkdir(parents=True, exist_ok=True)
    wb.save(SHEET)
    print(f"{len(sample)} items -> {SHEET}; key -> {KEY}")
    print(sample["man_d2"].value_counts().to_dict())


def yes_no(value) -> bool | None:
    v = str(value or "").strip().casefold()
    return True if v in YES else False if v in NO else None


def score(path) -> None:
    ws = load_workbook(path)["itens"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    human = {}
    for r in rows:
        t, la = yes_no(r[4]), yes_no(r[5])
        if t is None or la is None:
            raise SystemExit(f"item {r[0]}: fill both columns with sim/não")
        human[int(r[0])] = "both" if t and la else "technical" if t else "lay" if la else "none"
    key = pd.read_csv(KEY, dtype={"id": int})
    key["human_d2"] = key["id"].map(human)
    k = cohen_kappa(list(key["labeler_d2"]), list(key["human_d2"]))
    out = {
        "n_items": len(key),
        "per_class_drawn": PER_CLASS,
        "seed": SEED,
        "agreement": round(float((key["labeler_d2"] == key["human_d2"]).mean()), 4),
        "cohen_kappa": round(k, 4),
        "kappa_min": KAPPA_MIN,
        "d2_strata_status": "confirmatory" if k >= KAPPA_MIN else "exploratory",
        "confusion_labeler_x_human": pd.crosstab(key["labeler_d2"], key["human_d2"]).to_dict(),
        "checked_by": "project author (not a clinician)",
    }
    RESULT.write_text(json.dumps(out, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    key.to_csv(KEY, index=False)
    print(json.dumps(out, indent=2, ensure_ascii=False))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--make", action="store_true")
    parser.add_argument("--score", metavar="XLSX")
    args = parser.parse_args()
    if args.make:
        make()
    elif args.score:
        score(args.score)
    else:
        parser.print_help()


if __name__ == "__main__":
    main()
