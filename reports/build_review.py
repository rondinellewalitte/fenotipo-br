"""Build the review package for the most frequent HPO terms without an official PT label:
reports/review/top50_terms.{csv,xlsx} for a clinical geneticist, plus a README whose numbers
are read from results/. Optionally fills `candidate_pt` with machine translations.

    uv run python reports/build_review.py               # package only, no API calls
    uv run python reports/build_review.py --translate   # also call the Anthropic API

API calls happen only with --translate AND ANTHROPIC_API_KEY set; spending is capped at
BUDGET_USD (worst case checked before every call). Every translation is cached in
results/review/translations.jsonl and reused on later runs; calls are logged with tokens
and cost in results/review/llm_calls.jsonl.
"""
# ruff: noqa: E501  (PT-BR text templates below are prose)

from __future__ import annotations

import argparse
import json
import os
from datetime import UTC, datetime

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

from hpolab import llm
from hpolab.data import RAW_DIR, ROOT, load_manifest
from hpolab.hpo import Term, load_babelon, load_ontology, pt_labels
from hpolab.provenance import git_info
from hpolab.retrieval import text_key

N_TERMS = 50
BUDGET_USD = 2.00
MAX_TOKENS = 1500
EFFORT = "low"
PROMPT_STEM = ROOT / "prompts" / "translate_hpo_label_pt_br_v1"
MISSING = ROOT / "results" / "exp02" / "missing_terms.csv"
OOS = ROOT / "results" / "exp02" / "out_of_sample.json"
METRICS2 = ROOT / "results" / "exp02" / "metrics.json"
OUT = ROOT / "reports" / "review"
LOGS = ROOT / "results" / "review"
CACHE = LOGS / "translations.jsonl"
CALL_LOG = LOGS / "llm_calls.jsonl"

COLUMNS = [
    "hpo_id",
    "label_en",
    "definition_en",
    "synonyms_en",
    "n_cases",
    "n_genes",
    "candidate_pt",
    "reviewer_pt",
    "reviewer_notes",
]
TO_FILL = ("reviewer_pt", "reviewer_notes")
# Already officially translated, not in the list: used only to show the expected format.
EXAMPLE_IDS = ("HP:0000252", "HP:0004322")
SCHEMA = {
    "type": "object",
    "properties": {
        "candidate_pt": {"type": "string"},
        "alternatives": {"type": "array", "items": {"type": "string"}},
        "notes": {"type": "string"},
    },
    "required": ["candidate_pt", "alternatives", "notes"],
    "additionalProperties": False,
}
MT_WARNING = "TRADUÇÃO AUTOMÁTICA (gerada por máquina). Exige revisão humana; não usar sem revisão."
# Marker the reviewer writes in reviewer_notes when the candidate is kept verbatim; lets us
# measure machine-translation accuracy and counters anchoring on the candidate.
ACCEPTED_MARK = "candidata aceita sem alteração"
NO_FEAR = (
    "A tradução candidata é só um ponto de partida e pode estar errada ou soar estranha. "
    "Corrija sem receio sempre que preferir outro termo. Se aceitar a candidata exatamente "
    f"como está, copie-a em reviewer_pt e escreva em reviewer_notes: {ACCEPTED_MARK}. "
    "Isso permite medir quanto a tradução automática acertou."
)


def synonyms_text(term: Term) -> str:
    """Exact synonyms as-is; other scopes tagged, e.g. 'Epilepsy [related]'. Synonyms that
    only repeat the label are dropped."""
    label = text_key(term.label)
    parts = [
        s.text if s.scope == "exact" else f"{s.text} [{s.scope}]"
        for s in term.synonyms
        if text_key(s.text) != label
    ]
    return "; ".join(dict.fromkeys(parts))


def build_terms(ont) -> pd.DataFrame:
    top = pd.read_csv(MISSING).head(N_TERMS)
    rows = []
    for r in top.itertuples():
        t = ont.terms[r.hpo_id]
        rows.append(
            {
                "hpo_id": t.id,
                "label_en": t.label,
                "definition_en": t.definition,
                "synonyms_en": synonyms_text(t),
                "n_cases": int(r.n_cases),
                "n_genes": int(r.n_genes),
                "candidate_pt": "",
                "reviewer_pt": "",
                "reviewer_notes": "",
            }
        )
    return pd.DataFrame(rows, columns=COLUMNS)


def load_cache() -> dict[tuple, dict]:
    if not CACHE.exists():
        return {}
    out = {}
    for line in CACHE.read_text(encoding="utf-8").splitlines():
        rec = json.loads(line)
        out[(rec["hpo_id"], rec["label_en"], rec["model"], rec["prompt_sha256"])] = rec
    return out


def translate(df: pd.DataFrame, prompt: llm.Prompt, requested: bool) -> dict:
    budget = llm.Budget(BUDGET_USD)
    key_present = llm.has_api_key()
    s = {
        "api_key_present": key_present,
        "translate_requested": requested,
        "model": llm.MODEL,
        "prompt_files": [
            f"prompts/{PROMPT_STEM.name}.system.md",
            f"prompts/{PROMPT_STEM.name}.user.md",
        ],
        "prompt_sha256": prompt.sha256,
        "effort": EFFORT,
        "max_tokens": MAX_TOKENS,
        "sampling": "temperature/top_p/top_k are rejected by the API since 2026-10-21; "
        "reproducibility comes from the pinned model/prompt and the on-disk cache",
        "budget_usd": BUDGET_USD,
        "prices_usd_per_mtok": llm.PRICES[llm.MODEL],
        "n_terms": len(df),
        "n_from_cache": 0,
        "n_api_calls": 0,
        "n_not_finished": 0,
        "stopped_by_budget": None,
    }
    cache = load_cache()
    client = None
    for i, row in df.iterrows():
        key = (row.hpo_id, row.label_en, llm.MODEL, prompt.sha256)
        if key in cache:
            df.at[i, "candidate_pt"] = cache[key]["candidate_pt"]
            s["n_from_cache"] += 1
            continue
        if not (key_present and requested):
            continue
        if client is None:
            import anthropic

            client = anthropic.Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        user = prompt.render(
            hpo_id=row.hpo_id,
            label_en=row.label_en,
            definition_en=row.definition_en or "(none)",
            synonyms_en=row.synonyms_en or "(none)",
        )
        try:
            out = llm.json_call(
                client,
                model=llm.MODEL,
                prompt=prompt,
                user=user,
                schema=SCHEMA,
                max_tokens=MAX_TOKENS,
                effort=EFFORT,
                budget=budget,
                log_path=CALL_LOG,
                meta={"hpo_id": row.hpo_id},
            )
        except llm.BudgetExceeded as exc:
            s["stopped_by_budget"] = str(exc)
            break
        s["n_api_calls"] += 1
        if out is None:
            s["n_not_finished"] += 1
            continue
        llm.append_jsonl(
            CACHE,
            {
                "hpo_id": row.hpo_id,
                "label_en": row.label_en,
                "model": llm.MODEL,
                "prompt_sha256": prompt.sha256,
                "candidate_pt": out["candidate_pt"].strip(),
                "alternatives": out["alternatives"],
                "notes": out["notes"],
                "created_at": datetime.now(UTC).isoformat(timespec="seconds"),
            },
        )
        df.at[i, "candidate_pt"] = out["candidate_pt"].strip()
    s["spent_usd_this_run"] = round(budget.spent_usd, 4)
    calls = (
        [json.loads(x) for x in CALL_LOG.read_text(encoding="utf-8").splitlines()]
        if CALL_LOG.exists()
        else []
    )
    s["n_logged_api_calls"] = len(calls)
    s["logged_cost_usd_total"] = round(sum(c["cost_usd"] for c in calls), 4)
    s["n_with_candidate"] = int((df["candidate_pt"] != "").sum())
    avg_chars = len(prompt.system) + len(prompt.user_template) + 400
    s["worst_case_usd_for_all_terms"] = round(
        len(df) * budget.worst_case(llm.MODEL, avg_chars, MAX_TOKENS), 2
    )
    return s


def pct(x: float, d: int = 1) -> str:
    return f"{x * 100:.{d}f}%".replace(".", ",")


def pts(x: float, d: int = 1) -> str:
    return f"{'+' if x >= 0 else '−'}{abs(x) * 100:.{d}f}".replace(".", ",") + " p.p."


def provenance() -> dict:
    m = load_manifest()
    return {
        "hpo_release": m.get("hp.json", {}).get("release_tag"),
        "phenopacket_store_release": m.get("all_phenopackets.zip", {}).get("release_tag"),
        "generated_at": datetime.now(UTC).strftime("%Y-%m-%d"),
        "git": git_info(),
    }


def commit_label(prov: dict) -> str:
    sha = (prov["git"]["commit"] or "")[:7]
    return f"{sha}, com mudanças não commitadas" if prov["git"]["dirty"] else sha


def candidate_status(s: dict) -> str:
    if s["n_with_candidate"] == s["n_terms"]:
        return f"Preenchida para os {s['n_terms']} termos pelo modelo `{s['model']}` (prompt `{s['prompt_files'][0]}`, sha256 `{s['prompt_sha256'][:12]}`)."
    if s["n_with_candidate"] == 0:
        reason = (
            "a variável ANTHROPIC_API_KEY não estava definida"
            if not s["api_key_present"]
            else "o script rodou sem `--translate`"
        )
        return (
            f"**Vazia nesta versão**: {reason}. A revisão pode ser feita direto em `reviewer_pt`."
        )
    return f"Preenchida para {s['n_with_candidate']} de {s['n_terms']} termos (ver `results/review/run_summary.json`)."


def write_xlsx(df: pd.DataFrame, s: dict, prov: dict, examples: list[dict]) -> None:
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    fill_input = PatternFill("solid", fgColor="FFF2CC")  # cells the reviewer fills in
    fill_mt = PatternFill("solid", fgColor="FCE4D6")  # machine translation, read-only
    fill_head = PatternFill("solid", fgColor="D9D9D9")
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    readme = wb.active
    readme.title = "LEIA-ME"
    lines = [
        ("Revisão de termos HPO sem tradução oficial em PT-BR", None),
        ("", None),
        (
            "O que fazer",
            "Na aba 'termos', preencha só as colunas em amarelo: reviewer_pt (o rótulo final em PT-BR, mesmo que seja igual ao candidato) e reviewer_notes (opcional).",
        ),
        ("Corrija sem receio", NO_FEAR),
        (
            "candidate_pt",
            MT_WARNING
            + " Fundo laranja. "
            + candidate_status(s).replace("**", "").replace("`", ""),
        ),
        (
            "Convenções",
            "Grafia brasileira; termo usual na clínica e na genética médica no Brasil; forma de rótulo do HPO (substantivo, só a primeira letra maiúscula, sem ponto final).",
        ),
        (
            "label_en / definition_en / synonyms_en",
            "Rótulo, definição e sinônimos oficiais em inglês (HPO). Sinônimos que não são exatos aparecem marcados, por exemplo [related].",
        ),
        (
            "n_cases / n_genes",
            "Em quantos casos e genes do Phenopacket Store o termo aparece. A lista está em ordem decrescente de n_cases.",
        ),
        (
            "Crédito",
            "Nome e ORCID de quem revisar entram como tradutor(a) na submissão ao HPO e como coautor(a) no artigo.",
        ),
        ("Seu nome", ""),
        ("Seu ORCID", ""),
        ("", None),
        (
            "Exemplo de preenchimento (termos que já têm tradução oficial; não fazem parte da lista)",
            None,
        ),
    ]
    for r, (k, v) in enumerate(lines, start=1):
        readme.cell(r, 1, k).font = bold
        if v is not None:
            c = readme.cell(r, 2, v)
            c.font, c.alignment = base, wrap
        if k in ("Seu nome", "Seu ORCID"):
            readme.cell(r, 2).fill = fill_input
    readme.cell(1, 1).font = Font(name="Arial", size=14, bold=True)
    ex_head = ["hpo_id", "label_en", "reviewer_pt", "reviewer_notes"]
    r0 = len(lines) + 1
    for j, h in enumerate(ex_head, start=1):
        c = readme.cell(r0, j, h)
        c.font, c.fill = bold, fill_head
    for i, ex in enumerate(examples, start=r0 + 1):
        for j, h in enumerate(ex_head, start=1):
            readme.cell(i, j, ex[h]).font = base
    r1 = r0 + len(examples) + 2
    meta = [
        (
            "Fontes",
            f"HPO {prov['hpo_release']}; Phenopacket Store {prov['phenopacket_store_release']}; lista em results/exp02/missing_terms.csv",
        ),
        (
            "Gerado em",
            f"{prov['generated_at']} por reports/build_review.py (commit {commit_label(prov)})",
        ),
        ("Aviso", "Apoio à pesquisa; não é ferramenta diagnóstica nem orienta conduta clínica."),
    ]
    for k, (a, b) in enumerate(meta, start=r1):
        readme.cell(k, 1, a).font = bold
        c = readme.cell(k, 2, b)
        c.font, c.alignment = base, wrap
    readme.column_dimensions["A"].width = 34
    readme.column_dimensions["B"].width = 100
    readme.column_dimensions["C"].width = 28
    readme.column_dimensions["D"].width = 30

    ws = wb.create_sheet("termos")
    widths = {
        "hpo_id": 12,
        "label_en": 32,
        "definition_en": 60,
        "synonyms_en": 40,
        "n_cases": 9,
        "n_genes": 9,
        "candidate_pt": 32,
        "reviewer_pt": 32,
        "reviewer_notes": 40,
    }
    notes = {
        "candidate_pt": MT_WARNING,
        "reviewer_pt": "Preencha com o rótulo final em PT-BR (mesmo que igual ao candidato).",
        "reviewer_notes": f"Opcional: dúvidas, termo alternativo, uso regional. Se mantiver a candidata sem mudar nada, escreva: {ACCEPTED_MARK}",
    }
    for j, col in enumerate(COLUMNS, start=1):
        c = ws.cell(1, j, col)
        c.font, c.fill, c.alignment = bold, fill_head, wrap
        if col in notes:
            c.comment = Comment(notes[col], "hpolab")
        ws.column_dimensions[c.column_letter].width = widths[col]
    for i, row in enumerate(df[COLUMNS].itertuples(index=False), start=2):
        for j, col in enumerate(COLUMNS, start=1):
            c = ws.cell(i, j, getattr(row, col))
            c.font, c.alignment = base, wrap
            if col == "candidate_pt":
                c.fill, c.font = fill_mt, Font(name="Arial", size=10, italic=True)
            elif col in TO_FILL:
                c.fill = fill_input
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(OUT / "top50_terms.xlsx")


def write_readme(s: dict, prov: dict) -> None:
    oos, m2 = (
        json.loads(OOS.read_text(encoding="utf-8")),
        json.loads(METRICS2.read_text(encoding="utf-8")),
    )
    base = m2["coverage"]["annotation_coverage"]
    base_full = m2["coverage"]["pct_cases_full"]
    g = oos["gene_split"]["seeds_summary_annotation_coverage_gain"][f"top{N_TERMS}"]["oos_gain"]
    n_seeds = len(oos["gene_split"]["seeds"])
    ins = m2["simulated_gain"][f"top{N_TERMS}_missing_added"]
    hp = oos["hpoa_ranked_list"][f"top{N_TERMS}"]["hpoa_ranked"]
    a = oos["gene_split"]["primary"]["A_to_B"][f"top{N_TERMS}"]["out_of_sample"]["pct_cases_full"]
    b = oos["gene_split"]["primary"]["B_to_A"][f"top{N_TERMS}"]["out_of_sample"]["pct_cases_full"]
    md = f"""# Revisão de {N_TERMS} termos HPO sem tradução PT-BR

Gerado por `reports/build_review.py` em {prov["generated_at"]} (commit `{commit_label(prov)}`). Os números vêm de `results/`; não edite este arquivo à mão.

> Apoio à pesquisa. Não é ferramenta diagnóstica nem orienta conduta clínica (ANVISA RDC 657/2022).

## O que é

Os {N_TERMS} termos do HPO ({prov["hpo_release"]}) **sem rótulo oficial em português** que mais aparecem nos casos do Phenopacket Store ({prov["phenopacket_store_release"]}). A lista vem de `results/exp02/missing_terms.csv`.

- `top50_terms.xlsx`: para abrir no Excel. A aba **LEIA-ME** tem as instruções e um exemplo; a aba **termos** tem a lista.
- `top50_terms.csv`: o mesmo conteúdo em texto (UTF-8).

## Por que estes termos

- **Hoje:** {pct(base)} das anotações de fenótipo desses casos têm rótulo PT oficial, e só {pct(base_full)} dos casos estão 100% traduzidos (`results/exp02/metrics.json`).
- **Ganho esperado** com a tradução destes {N_TERMS} termos: **≈{pct(base, 0)} → ≈{pct(base + g["mean"], 0)} de cobertura por anotação**, fora da amostra (`results/exp02/out_of_sample.json`).
  - O ganho médio é de {pts(g["mean"])} (faixa {pts(g["min"])} a {pts(g["max"])}). Ele foi medido em {n_seeds} divisões dos genes em duas metades: a lista é ordenada numa metade e o ganho é medido na outra.
  - Casos 100% traduzidos: de {pct(base_full)} para {pct(min(a, b))}–{pct(max(a, b))}, fora da amostra.
  - Para comparação: dentro da amostra, que é um teto otimista, a cobertura chega a {pct(ins["annotation_coverage"])}. Com uma lista ordenada pelo `phenotype.hpoa`, que independe destes casos, chega a {pct(hp["annotation_coverage"])}.
- **Limites:**
  - Os casos vêm da literatura e foram anotados em inglês. A frequência num ambulatório brasileiro pode ser outra.
  - **Viés de ancoragem:** quem revisa tende a aprovar a tradução candidata que já está na planilha. Por isso pedimos que corrija sem receio e marque explicitamente quando aceitar a candidata sem alteração. A taxa de aceitação deve ser lida como teto do acerto da tradução automática, não como medida independente.

## Como revisar

Preencha as duas colunas em amarelo:

- **`reviewer_pt`:** o rótulo final em PT-BR, mesmo que seja igual ao candidato. Use grafia brasileira, o termo usual na clínica e na genética médica no Brasil, e a forma de rótulo do HPO (substantivo, só a primeira letra maiúscula, sem ponto final).
- **`reviewer_notes`:** opcional. Dúvidas, termos alternativos, uso regional.
- **Corrija sem receio.** {NO_FEAR}

As colunas `label_en`, `definition_en` e `synonyms_en` trazem o conceito oficial em inglês. Sinônimos que não são exatos vêm marcados, por exemplo `[related]`.

## Sobre `candidate_pt`

**{MT_WARNING}**

{candidate_status(s)}

- **Configuração:** modelo fixo `{s["model"]}`, esforço `{s["effort"]}` e saída em JSON validado por esquema.
  - A API não aceita mais o parâmetro de temperatura. Por isso, cada resposta é guardada em `results/review/translations.jsonl` e reaproveitada, em vez de ser gerada de novo.
- **Custo:** as chamadas, com tokens e custo, ficam em `results/review/llm_calls.jsonl`. O resumo está em `results/review/run_summary.json`, com teto de US$ {f"{s['budget_usd']:.2f}".replace(".", ",")}. Total gasto até agora: US$ {f"{s['logged_cost_usd_total']:.2f}".replace(".", ",")} em {s["n_logged_api_calls"]} chamadas.

## Depois da revisão

1. **Conversão:** o rótulo de `reviewer_pt` vira uma linha no formato babelon do HPO (`src/translations/hp-pt.babelon.tsv`), com o ORCID de quem revisou na coluna `translator`.
2. **Submissão:** um pull request ao repositório `obophenotype/human-phenotype-ontology`, com crédito de tradução e coautoria no artigo do piloto.
"""
    (OUT / "README.md").write_text(md, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--translate", action="store_true", help="call the API for candidate_pt")
    args = parser.parse_args()
    ont = load_ontology(RAW_DIR / "hp.json")
    df = build_terms(ont)
    prompt = llm.Prompt.load(PROMPT_STEM)
    s = translate(df, prompt, args.translate)
    pt = pt_labels(load_babelon(RAW_DIR / "hp-pt.babelon.tsv"), ont).set_index("hpo_id")
    examples = [
        {
            "hpo_id": i,
            "label_en": ont.terms[i].label,
            "reviewer_pt": pt.loc[i, "pt_label"],
            "reviewer_notes": "",
        }
        for i in EXAMPLE_IDS
        if i in pt.index and i not in set(df["hpo_id"])
    ]
    prov = provenance()
    OUT.mkdir(parents=True, exist_ok=True)
    LOGS.mkdir(parents=True, exist_ok=True)
    df[COLUMNS].to_csv(OUT / "top50_terms.csv", index=False, encoding="utf-8")
    write_xlsx(df, s, prov, examples)
    write_readme(s, prov)
    s.update({"generated_at": prov["generated_at"], "git": prov["git"]})
    (LOGS / "run_summary.json").write_text(
        json.dumps(s, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(
        json.dumps(
            {
                k: s[k]
                for k in (
                    "api_key_present",
                    "translate_requested",
                    "n_with_candidate",
                    "n_api_calls",
                    "spent_usd_this_run",
                    "worst_case_usd_for_all_terms",
                )
            },
            indent=2,
        )
    )
    if not s["api_key_present"]:
        print("WARNING: ANTHROPIC_API_KEY not set; candidate_pt left empty.")


if __name__ == "__main__":
    main()
